from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
EVIDENCE = DOCS / "evidencias"
OUTPUT = DOCS / "MR222768_Desafio03_Entrega.docx"
REPOSITORY_URL = "https://github.com/Martinez014/MR222768_desafio03"

NAVY = "1F4E78"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F5F6F7"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), BORDER)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for i, text in enumerate(headers):
        cell = header_row.cells[i]
        cell.text = text
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, LIGHT_BLUE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(color)
    run_properties.append(underline)
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])


def set_alt_text(inline_shape, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_figure(doc: Document, image_path: Path, width_inches: float, caption: str, alt: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = paragraph.add_run().add_picture(str(image_path), width=Inches(width_inches))
    set_alt_text(shape, caption, alt)
    caption_paragraph = doc.add_paragraph(caption, style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(8)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(text, style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = False


def add_continuation_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(15)
    run.font.color.rgb = BLACK


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_after = Pt(12)
    title_p_pr = title.element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for name, size in (("Heading 1", 15), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = False

    caption = styles["Caption"]
    caption.font.name = "Arial"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = BLACK

    footer = section.footer.paragraphs[0]
    footer.text = "MR222768_desafío03  |  Página "
    footer.style = styles["Footer"]
    add_page_number(footer)
    for run in footer.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(90, 90, 90)


def read_results() -> tuple[dict, int]:
    http_results = json.loads((EVIDENCE / "06-http-integration.json").read_text(encoding="utf-8-sig"))
    test_text = (EVIDENCE / "03-dotnet-test.txt").read_text(encoding="utf-8-sig")
    match = re.search(r"Pruebas totales:\s*(\d+)", test_text)
    total_tests = int(match.group(1)) if match else 19
    return http_results, total_tests


def build_document() -> None:
    http, total_tests = read_results()
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Sistema de Pedidos con ASP.NET Core Web API"
    doc.core_properties.subject = "Documento de entrega del Desafío 03"
    doc.core_properties.author = "MR222768"
    doc.core_properties.keywords = "ASP.NET Core, Web API, Docker, SQL Server, xUnit"

    # Portada
    doc.add_paragraph().paragraph_format.space_after = Pt(65)
    institution = doc.add_paragraph("Universidad Don Bosco")
    institution.alignment = WD_ALIGN_PARAGRAPH.CENTER
    institution.runs[0].bold = True
    institution.runs[0].font.size = Pt(16)
    institution.paragraph_format.space_after = Pt(35)

    challenge = doc.add_paragraph("Desafío 03")
    challenge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    challenge.runs[0].font.size = Pt(14)
    challenge.paragraph_format.space_after = Pt(18)

    title = doc.add_paragraph("Sistema de Pedidos con ASP.NET Core Web API", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p_pr = title._p.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    solution = doc.add_paragraph("Solución: MR222768_desafío03")
    solution.alignment = WD_ALIGN_PARAGRAPH.CENTER
    solution.runs[0].font.size = Pt(13)
    solution.paragraph_format.space_after = Pt(8)

    cif = doc.add_paragraph("CIF: MR222768")
    cif.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cif.runs[0].font.size = Pt(12)
    doc.add_page_break()

    add_heading(doc, "1 Introducción")
    doc.add_paragraph(
        "Este documento presenta el sistema de pedidos desarrollado para el Desafío 03. La aplicación permite administrar clientes, productos y pedidos mediante una API REST. El proyecto utiliza una arquitectura de cuatro capas, persistencia con Entity Framework Core y SQL Server, ejecución con Docker y pruebas unitarias con xUnit."
    )
    doc.add_paragraph(
        "Las comprobaciones incluidas corresponden a ejecuciones reales del proyecto: compilación correcta, 19 pruebas unitarias aprobadas y un recorrido HTTP completo sobre los contenedores de la API y la base de datos."
    )

    add_heading(doc, "2 Objetivo del proyecto")
    doc.add_paragraph(
        "Desarrollar una API mantenible para registrar clientes y productos, crear pedidos y controlar el inventario. La solución debe validar los datos de entrada, aplicar las reglas de negocio antes de guardar cambios y ofrecer respuestas HTTP claras para cada resultado."
    )

    add_heading(doc, "3 Tecnologías utilizadas")
    add_table(
        doc,
        ["Tecnología", "Uso en el proyecto"],
        [
            [".NET 8 y C#", "Plataforma y lenguaje de la solución."],
            ["ASP.NET Core Web API", "Controladores, rutas REST, validación automática y OpenAPI."],
            ["Entity Framework Core", "Acceso a datos Code First y migraciones."],
            ["SQL Server 2022", "Base de datos relacional del sistema."],
            ["Docker Compose", "Ejecución coordinada de la API y SQL Server."],
            ["xUnit", "Pruebas unitarias de servicios, validaciones y reglas de negocio."],
        ],
        [1.75, 4.95],
    )

    add_heading(doc, "4 Arquitectura de N capas")
    doc.add_paragraph(
        "La solución separa las responsabilidades para evitar que los controladores contengan acceso a datos o lógica de negocio. Las dependencias apuntan hacia las capas internas."
    )
    add_table(
        doc,
        ["Capa", "Responsabilidad"],
        [
            ["Domain", "Entidades Cliente, Producto y Pedido, junto con los conceptos centrales del dominio."],
            ["Application", "DTO, interfaces, servicios y reglas de negocio que coordinan las operaciones."],
            ["Infrastructure", "DbContext, repositorios, configuración de EF Core, SQL Server y migraciones."],
            ["API", "Controladores REST, inyección de dependencias, OpenAPI y manejo uniforme de errores."],
            ["Tests", "Pruebas unitarias aisladas con repositorios de prueba, sin depender de SQL Server ni Docker."],
        ],
        [1.55, 5.15],
    )
    doc.add_page_break()

    add_heading(doc, "5 Modelos Clientes Productos y Pedidos")
    doc.add_paragraph(
        "Los modelos cumplen los campos solicitados y utilizan Data Annotations en las entidades y en los DTO de entrada. EF Core configura de forma explícita las claves foráneas, los índices y la precisión decimal del precio."
    )
    add_table(
        doc,
        ["Modelo", "Campos principales", "Validaciones relevantes"],
        [
            ["Cliente", "Id, Nombre, Apellido, Email", "Nombre y apellido obligatorios; longitudes máximas; correo obligatorio y válido."],
            ["Producto", "Id, Nombre, Descripcion, Precio, Stock", "Nombre obligatorio; precio mayor que cero; stock no negativo; precisión decimal para moneda."],
            ["Pedido", "Id, FechaPedido, Cantidad, ClienteId, ProductoId", "Cantidad mayor que cero; fecha válida; cliente y producto existentes."],
        ],
        [1.1, 2.7, 2.9],
    )
    doc.add_paragraph(
        "Cada pedido se relaciona con un cliente y un producto. Los DTO de consulta incluyen la información necesaria para identificar ambas entidades sin serializar las propiedades de navegación y sin producir ciclos JSON."
    )

    add_heading(doc, "6 Funcionalidades de la API REST")
    doc.add_paragraph(
        "Los tres recursos disponen de operaciones CRUD reales. Las respuestas usan 200 para consultas y actualizaciones, 201 para creaciones, 204 para eliminaciones, 400 para datos inválidos, 404 para registros inexistentes y 409 para conflictos de negocio."
    )
    add_table(
        doc,
        ["Recurso", "Operaciones disponibles"],
        [
            ["Clientes", "GET /api/clientes, GET /api/clientes/{id}, POST, PUT y DELETE."],
            ["Productos", "GET /api/productos, GET /api/productos/{id}, POST, PUT y DELETE."],
            ["Pedidos", "GET /api/pedidos, GET /api/pedidos/{id}, POST, PUT y DELETE."],
        ],
        [1.25, 5.45],
    )
    doc.add_paragraph(
        "Swagger UI publica el contrato OpenAPI y permite revisar los endpoints desde el navegador en http://localhost:8080/swagger."
    )

    add_heading(doc, "7 Validaciones y reglas de negocio")
    for item in (
        "No se aceptan clientes con campos obligatorios vacíos ni correos con formato inválido.",
        "El precio debe ser mayor que cero y el stock no puede ser negativo.",
        "Un pedido requiere cliente y producto existentes, cantidad positiva y stock suficiente.",
        "Al crear un pedido se descuenta el inventario; al actualizarlo se ajusta la diferencia y se manejan cambios de producto.",
        "Al eliminar un pedido se devuelve al inventario la cantidad correspondiente.",
        "No se pueden eliminar clientes o productos asociados a pedidos; la API devuelve 409 Conflict.",
        "Los cambios del pedido y del inventario se guardan de forma atómica.",
        "El inventario usa una columna RowVersion como token de concurrencia. Los conflictos de escritura se controlan para impedir ventas por encima del stock disponible.",
    ):
        add_bullet(doc, item)
    doc.add_page_break()

    add_heading(doc, "8 Ejecución con Docker")
    doc.add_paragraph(
        "El Dockerfile usa compilación multi-stage y publica únicamente la salida necesaria de la API. Docker Compose inicia SQL Server 2022 y la API en una red interna, conserva los datos en un volumen y expone la aplicación por el puerto 8080. La contraseña de SQL Server se obtiene del archivo .env local, que está excluido del repositorio."
    )
    doc.add_paragraph(
        "El contenedor de SQL Server incluye una comprobación de salud. La API espera la disponibilidad de la base de datos y aplica las migraciones de EF Core al iniciar, con reintentos para el primer arranque."
    )
    add_table(
        doc,
        ["Servicio", "Estado comprobado", "Puerto local"],
        [
            ["sistema-pedidos-sqlserver-1", "En ejecución y saludable", "1433"],
            ["sistema-pedidos-api-1", "En ejecución", "8080"],
        ],
        [3.0, 2.4, 1.3],
    )
    doc.add_paragraph(
        "La migración inicial 20261017032521_InitialCreate figura aplicada en la tabla __EFMigrationsHistory de SQL Server."
    )

    add_heading(doc, "9 Pruebas unitarias xUnit y resultados")
    doc.add_paragraph(
        f"La ejecución de dotnet test terminó correctamente con {total_tests} pruebas aprobadas, sin fallos ni pruebas omitidas. Las pruebas unitarias no requieren contenedores ni una instancia de SQL Server."
    )
    add_table(
        doc,
        ["Grupo de prueba", "Cobertura comprobada"],
        [
            ["Clientes", "Creación válida, campos obligatorios, correo inválido y consulta inexistente."],
            ["Productos", "Creación válida, precio inválido, stock negativo y restricciones de eliminación."],
            ["Pedidos", "Creación, consulta, stock insuficiente, descuento, actualización, cambio de producto y restauración al eliminar."],
            ["Resultado", f"{total_tests} de {total_tests} pruebas aprobadas."],
        ],
        [1.55, 5.15],
    )
    doc.add_paragraph(
        "La compilación completa también terminó con 0 advertencias y 0 errores."
    )
    doc.add_page_break()

    add_heading(doc, "10 Evidencias del funcionamiento")
    doc.add_paragraph(
        "Las siguientes capturas son archivos independientes insertados en Word. Se pueden seleccionar, eliminar o sustituir sin modificar el resto del documento."
    )
    add_figure(
        doc,
        EVIDENCE / "09-docker-desktop.png",
        5.1,
        "Figura 1. Docker Desktop con el conjunto de contenedores del sistema en ejecución.",
        "Captura real de Docker Desktop que muestra el grupo sistema-pedido en ejecución.",
    )
    doc.add_page_break()
    add_continuation_heading(doc, "10 Evidencias del funcionamiento")
    add_figure(
        doc,
        EVIDENCE / "10-swagger.png",
        6.25,
        "Figura 2. Swagger UI con los endpoints de clientes, pedidos y productos.",
        "Captura real de Swagger UI que muestra los endpoints REST del Sistema de Pedidos API.",
    )
    doc.add_page_break()

    add_continuation_heading(doc, "10 Evidencias del funcionamiento")
    doc.add_paragraph(
        "La prueba HTTP automatizada recorrió las operaciones principales contra la API ejecutada en Docker. La tabla resume los códigos y cambios de inventario guardados en el archivo de evidencia."
    )
    add_table(
        doc,
        ["Comprobación HTTP", "Resultado"],
        [
            ["Swagger", str(http["swagger"])],
            ["Crear, consultar y actualizar cliente", f'{http["crearCliente"]}, {http["consultarCliente"]}, {http["actualizarCliente"]}'],
            ["Crear, consultar y actualizar producto", f'{http["crearProducto"]}, {http["consultarProducto"]}, {http["actualizarProducto"]}'],
            ["Crear, consultar y actualizar pedido", f'{http["crearPedido"]}, {http["consultarPedido"]}, {http["actualizarPedido"]}'],
            ["Datos inválidos", str(http["validacionIncorrecta"])],
            ["Stock insuficiente", str(http["stockInsuficiente"])],
            ["Registro inexistente", str(http["registroInexistente"])],
            ["Eliminar entidades con pedido asociado", f'{http["eliminarClienteConPedido"]} y {http["eliminarProductoConPedido"]}'],
            ["Eliminar pedido", str(http["eliminarPedido"])],
        ],
        [4.75, 1.95],
    )
    add_table(
        doc,
        ["Comprobación de inventario", "Unidades"],
        [
            ["Stock inicial", str(http["stockInicial"])],
            ["Después de crear un pedido de 3 unidades", str(http["stockDespuesCrear"])],
            ["Después de actualizar el pedido a 5 unidades", str(http["stockDespuesActualizar"])],
            ["Después de eliminar el pedido", str(http["stockRestaurado"])],
        ],
        [5.4, 1.3],
    )
    doc.add_paragraph(
        "Los datos usados por el recorrido HTTP se eliminaron al finalizar. La base de datos y el volumen de Docker se conservaron."
    )

    add_heading(doc, "11 Enlace al repositorio GitHub")
    doc.add_paragraph(
        "El código fuente, las migraciones, las pruebas, la configuración de Docker, las evidencias y este documento se encuentran en el repositorio privado:"
    )
    link_paragraph = doc.add_paragraph()
    link_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_hyperlink(link_paragraph, REPOSITORY_URL, REPOSITORY_URL)
    doc.add_paragraph(
        "El archivo .env no está versionado. Para clonar el repositorio es necesario utilizar una cuenta de GitHub con acceso autorizado."
    )

    add_heading(doc, "12 Conclusión")
    doc.add_paragraph(
        "El proyecto cumple el alcance solicitado para el sistema de pedidos. La API implementa el CRUD de clientes, productos y pedidos, valida la entrada, protege las reglas de inventario y persiste la información mediante EF Core y SQL Server. Las pruebas unitarias y el recorrido HTTP confirmaron el comportamiento principal, incluidos los conflictos por stock y la restauración del inventario."
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()

